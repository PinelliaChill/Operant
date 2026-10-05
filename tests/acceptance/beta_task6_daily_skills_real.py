"""H-12: six bundled Skills through formal commands and real model runs.

Runs only in a new absolute synthetic workspace/database. Evidence excludes
prompt/response bodies and credentials. The model must issue formal workspace
tool calls for file-producing Skills; file signatures and content are read back.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import sys
import time
from pathlib import Path

import httpx

from operant.api import create_app
from operant.application.default_skill_pack import DEFAULT_SKILL_PACK, install_default_skill_pack
from operant.application.security import PolicyEngine
from operant.contracts.b2_3 import ManagementCommand
from operant.domain.models import (
    Budget,
    CommandExecutionPolicy,
    CommandRunnerType,
    Effort,
    ModelProfile,
    RolePreset,
    ToolPolicy,
)
from operant.domain.security import PolicyBundle, PolicyDecision, PolicyLayer, PolicyRule
from operant.domain.threads import ConversationThread
from operant.settings import load_local_env

MODEL = "gpt-6-luna"
SPEC = {
    "documents": (
        "doc.json",
        {
            "title": "Daily brief",
            "sections": [
                {
                    "heading": "Result",
                    "paragraphs": [
                        "Operant created this document through its formal Skill command."
                    ],
                    "bullets": ["Read back the actual file."],
                }
            ],
        },
        "brief.docx",
        "docx",
    ),
    "presentations": (
        "deck.json",
        {
            "title": "Daily review",
            "slides": [
                {"title": "Result", "bullets": ["Formal Skill command", "Verified PowerPoint file"]}
            ],
        },
        "review.pptx",
        "pptx",
    ),
    "pdf": (
        "pdf.json",
        {
            "title": "Daily PDF",
            "sections": [
                {
                    "heading": "Result",
                    "paragraphs": ["Operant created and reopened this PDF."],
                    "bullets": [],
                }
            ],
        },
        "brief.pdf",
        "pdf",
    ),
    "skill-creator": (
        "skill.json",
        {
            "name": "daily-review",
            "description": "Review a daily note for unsupported claims.",
            "instructions": "Compare each claim with the supplied source and identify any gap.",
        },
        "daily-review/SKILL.md",
        "skill",
    ),
}


def _policy() -> PolicyEngine:
    return PolicyEngine(
        PolicyBundle(
            bundle_id="task6-daily-skills",
            version="test.v1",
            default_decision=PolicyDecision.DENY,
            rules=(
                PolicyRule(
                    rule_id="isolated-workspace",
                    layer=PolicyLayer.SYSTEM,
                    decision=PolicyDecision.ALLOW,
                    reason="synthetic acceptance workspace",
                ),
            ),
        )
    )


def _validate(name: str, workspace: Path, answer: str) -> dict[str, object]:
    if name == "grill-me":
        return {"question_present": "?" in answer or "？" in answer}
    if name == "find-skills":
        return {"catalog_match_present": "pdf" in answer.casefold()}
    _, _, output_name, kind = SPEC[name]
    output = workspace / output_name
    if not output.is_file() or output.stat().st_size < 100:
        return {"file_exists": False}
    if kind == "skill":
        from operant.skills.discovery import SkillDiscovery

        found = SkillDiscovery((workspace,)).discover()
        valid = any(item.name == "daily-review" for item in found.candidates)
        return {"file_exists": True, "valid_skill": valid, "bytes": output.stat().st_size}
    if kind == "docx":
        from docx import Document

        doc = Document(output)
        text = "\n".join(item.text for item in doc.paragraphs)
        valid = "Daily brief" in text and "Operant created" in text
    elif kind == "pptx":
        from pptx import Presentation

        deck = Presentation(output)
        valid = len(deck.slides) == 2 and deck.slides[0].shapes.title.text == "Daily review"
    else:
        from pypdf import PdfReader

        reader = PdfReader(output)
        valid = len(reader.pages) >= 1 and output.read_bytes().startswith(b"%PDF")
    return {"file_exists": True, "readback_valid": valid, "bytes": output.stat().st_size}


async def main(evidence_dir: Path, env_file: Path | None, selected_names: set[str] | None) -> None:
    if not evidence_dir.is_absolute():
        raise ValueError("evidence directory must be absolute")
    evidence_dir.mkdir(parents=True, exist_ok=False)
    workspace = evidence_dir / "workspace"
    workspace.mkdir()
    for _, (filename, spec, _, _) in SPEC.items():
        (workspace / filename).write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    if env_file is not None:
        load_local_env(env_file)
    base_url = os.environ["OPERANT_BASE_URL"]
    if not os.environ.get("OPERANT_API_KEY"):
        raise RuntimeError("OPERANT_API_KEY is unavailable")
    evidence: dict[str, object] = {
        "status": "failed",
        "model_id": MODEL,
        "entry": "executeSkillCommand -> ApplicationService.run_session -> WorkspaceTools",
        "synthetic_workspace": str(workspace),
        "cases": {},
    }
    app = create_app(
        evidence_dir / "core.sqlite3",
        artifact_root=evidence_dir / "artifacts",
        phase45_policy_engine=_policy(),
        phase56_local_authorizer=lambda _request: True,
    )
    service = app.state.operant_service
    manager = service.memory_manager_factory()
    try:
        discovered = await service.discover_models(base_url=base_url, secret_ref="OPERANT_API_KEY")
        if MODEL not in discovered:
            raise RuntimeError("formal Discovery did not return exact gpt-6-luna")
        evidence["discovery_exact_match"] = True
        project = await manager.execute(
            ManagementCommand(
                action="project_create", name="task6-daily-skills", workspace_path=str(workspace)
            )
        )
        ids = await install_default_skill_pack(
            manager, project_id=project.state.projects[-1].project_id
        )
        names = {
            entry.skill_name: skill_id
            for entry, skill_id in zip(DEFAULT_SKILL_PACK, ids, strict=True)
        }
        evidence["installed_enabled_count"] = len(ids)
        profile = service.add_model_profile(
            ModelProfile(
                name="task6-daily-skills",
                model_id=MODEL,
                base_url=base_url,
                secret_ref="OPERANT_API_KEY",
                supported_efforts=(Effort.LOW,),
                default_effort=Effort.LOW,
                effort_parameter=None,
            )
        )
        role = service.create_role(
            RolePreset(
                name="task6-daily-skills",
                system_prompt=(
                    "Use the explicitly selected Skill. For requested files, call "
                    "run_command with the supplied runtime Python and spec; do not "
                    "merely describe a file. Give a short result."
                ),
                model_profile_id=profile.id,
                effort=Effort.LOW,
                tool_policy=ToolPolicy(
                    allowed_tools=("run_command",),
                    workspace_write=True,
                    command_execution=True,
                    approval_required=(),
                    command_execution_policy=CommandExecutionPolicy(runner=CommandRunnerType.HOST),
                ),
                budget=Budget(
                    max_turns=8, max_tool_calls=4, timeout_seconds=180, max_output_tokens=800
                ),
                memory_scope="none",
            )
        )
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport, base_url="http://testserver", timeout=240
        ) as client:
            for index, name in enumerate(names):
                if selected_names is not None and name not in selected_names:
                    continue
                thread = service.create_thread(
                    ConversationThread(workspace_ref=str(workspace.resolve()))
                )
                session = service.create_session(role.id, thread_id=thread.id)
                if name in SPEC:
                    filename, _, output, kind = SPEC[name]
                    action = "create-skill" if kind == "skill" else f"make {kind}"
                    prompt = (
                        f"Use this Skill now. The JSON input already exists at {filename}. "
                        f"Run {sys.executable} -m operant.default_skill_tools {action} "
                        f"{filename} {output if kind != 'skill' else 'daily-review'}. "
                        "Use run_command with argv strings and cwd '.'. Report the command result."
                    )
                elif name == "find-skills":
                    prompt = (
                        f"Use run_command to invoke {sys.executable} -m "
                        "operant.default_skill_tools find-skills pdf. Report the real match."
                    )
                else:
                    prompt = (
                        "Begin grilling my plan to create a three-page project brief. "
                        "Ask one focused question about the most important missing "
                        "requirement. Do not answer it yourself."
                    )
                body = {
                    "command": f"skill:{names[name]}",
                    "arguments": {"prompt": prompt},
                    "idempotency_key": f"task6-{index}-{name}",
                }
                pending = asyncio.create_task(
                    client.post(
                        f"/v1/workbench/threads/{thread.id}/skill-commands",
                        json=body,
                        headers={"Idempotency-Key": body["idempotency_key"]},
                    )
                )
                approved_calls: set[str] = set()
                deadline = time.monotonic() + 170
                while not pending.done() and time.monotonic() < deadline:
                    for approval in service.list_pending_approvals(session.id):
                        call_id = str(approval["tool_call_id"])
                        if call_id in approved_calls:
                            continue
                        detail = str(approval["detail"])
                        if (
                            name not in SPEC
                            and name != "find-skills"
                            or approval["category"] != "security_policy"
                            or "executable=python" not in detail
                        ):
                            raise AssertionError("unexpected tool approval request")
                        decision = service.decide_approval(
                            session.id,
                            call_id,
                            approved=True,
                            decided_by="user",
                            reason_code="isolated_acceptance_python",
                        )
                        if not decision["accepted"]:
                            raise AssertionError("isolated Python approval was rejected")
                        approved_calls.add(call_id)
                    await asyncio.sleep(0.1)
                response = await pending
                data = response.json()
                answer = str(data.get("result", ""))
                readback = _validate(name, workspace, answer)
                events = service.list_events(session.id)
                case = {
                    "http_status": response.status_code,
                    "command_status": data.get("status"),
                    "session_completed": any(
                        event.event_type == "agent.completed" for event in events
                    ),
                    "tool_completed": any(
                        event.event_type == "tool.completed"
                        and event.payload.get("name") == "run_command"
                        for event in events
                    ),
                    "manual_approvals": len(approved_calls),
                    **readback,
                }
                evidence["cases"][name] = case  # type: ignore[index]
                if response.status_code != 200 or data.get("status") != "completed":
                    raise AssertionError(f"{name} formal Skill command failed")
                if name == "grill-me" and not case["question_present"]:
                    raise AssertionError("grill-me did not ask a question")
                if name == "find-skills" and not case["catalog_match_present"]:
                    raise AssertionError("find-skills did not return a match")
                if name != "grill-me" and not case["tool_completed"]:
                    raise AssertionError(f"{name} did not use the formal workspace tool")
                if name in SPEC and not (case.get("readback_valid") or case.get("valid_skill")):
                    raise AssertionError(f"{name} did not produce a valid file")
                if name == "grill-me":
                    followup = await client.post(
                        f"/v1/workbench/threads/{thread.id}/skill-commands",
                        json={
                            "command": body["command"],
                            "arguments": {
                                "prompt": (
                                    "My audience is new teammates. Continue the interview "
                                    "from my answer: ask the next most important question "
                                    "and keep the confirmed audience in mind."
                                )
                            },
                            "idempotency_key": "task6-grill-followup",
                        },
                        headers={"Idempotency-Key": "task6-grill-followup"},
                    )
                    continued = str(followup.json().get("result", ""))
                    completed_turns = [
                        event
                        for event in service.list_events(session.id)
                        if event.event_type == "agent.completed"
                    ]
                    case["followup_session_completed"] = len(completed_turns) >= 2
                    case["followup_question_distinct"] = (
                        "?" in continued or "？" in continued
                    ) and continued.strip() != answer.strip()
                    case["answer_carried_forward"] = (
                        "new teammates" in continued.casefold() or "新同事" in continued
                    )
                    if not (
                        followup.status_code == 200
                        and followup.json().get("status") == "completed"
                        and case["followup_session_completed"]
                        and case["followup_question_distinct"]
                        and case["answer_carried_forward"]
                    ):
                        raise AssertionError("grill-me did not continue after the user answer")
        evidence["status"] = "passed"
    except Exception as exc:
        evidence["error_type"] = type(exc).__name__
        raise
    finally:
        (evidence_dir / "result.json").write_text(
            json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        await manager.close()
        service.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--evidence-dir", required=True, type=Path)
    parser.add_argument("--env-file", type=Path)
    parser.add_argument("--skills", help="Comma-separated Skill names to run; defaults to all six")
    args = parser.parse_args()
    selected = set(args.skills.split(",")) if args.skills else None
    if selected is not None and selected - {entry.skill_name for entry in DEFAULT_SKILL_PACK}:
        raise ValueError("unknown Skill selection")
    asyncio.run(main(args.evidence_dir, args.env_file, selected))
